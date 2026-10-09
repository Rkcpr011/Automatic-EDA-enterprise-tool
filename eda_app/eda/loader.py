"""Dataset loading: CSV (encoding/delimiter detection) and Excel (sheet discovery)."""
import codecs
import csv
import warnings

import pandas as pd

_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


def _detect_encoding(path):
    with open(path, "rb") as fh:
        chunk = fh.read(1_000_000)
    for enc in _ENCODINGS[:2]:
        try:
            # incremental decoder tolerates a multibyte char cut at the chunk boundary
            codecs.getincrementaldecoder(enc)().decode(chunk, final=False)
            return enc
        except UnicodeDecodeError:
            continue
    return "cp1252" if _try(chunk, "cp1252") else "latin-1"


def _try(chunk, enc):
    try:
        chunk.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def _detect_delimiter(path, encoding):
    with open(path, "r", encoding=encoding, errors="replace") as fh:
        sample = fh.read(100_000)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def list_sheets(path, ext):
    """Return [{'name', 'rows', 'cols'}] (dimensions are best-effort)."""
    try:
        return _list_sheets(path, ext)
    except Exception as exc:
        raise ValueError(f"Could not open the workbook (is it corrupt or password-protected?): {exc}")


def _list_sheets(path, ext):
    sheets = []
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                sheets.append({"name": ws.title, "rows": ws.max_row, "cols": ws.max_column})
        finally:
            wb.close()
    else:  # xls
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            for name in book.sheet_names():
                sh = book.sheet_by_name(name)
                sheets.append({"name": name, "rows": sh.nrows, "cols": sh.ncols})
        finally:
            book.release_resources()
    return sheets


def _normalize(df):
    """Tidy column names; drop fully-empty rows and empty 'Unnamed' columns."""
    cols, seen = [], {}
    for i, c in enumerate(df.columns):
        name = str(c).strip() or f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        cols.append(name)
    df.columns = cols
    empty_unnamed = [c for c in df.columns if c.startswith("Unnamed") and df[c].isna().all()]
    df = df.drop(columns=empty_unnamed)
    df = df.dropna(how="all").reset_index(drop=True)
    return df


def load_dataset(path, ext, sheet=None):
    """Load a dataset into a DataFrame. Raises ValueError with a friendly message."""
    path = str(path)
    try:
        if ext == "csv":
            enc = _detect_encoding(path)
            sep = _detect_delimiter(path, enc)
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                df = pd.read_csv(path, sep=sep, encoding=enc, low_memory=False,
                                 on_bad_lines="skip")
        elif ext in ("xlsx", "xls"):
            engine = "openpyxl" if ext == "xlsx" else "xlrd"
            df = pd.read_excel(path, sheet_name=sheet if sheet is not None else 0, engine=engine)
        else:
            raise ValueError(f"Unsupported file type: .{ext}")
    except pd.errors.EmptyDataError:
        raise ValueError("The file is empty or has no readable columns.")
    except ValueError:
        raise
    except Exception as exc:  # corrupt file, wrong format, password protected ...
        raise ValueError(f"Could not read the file: {exc}")

    df = _normalize(df)
    if df.shape[0] == 0 or df.shape[1] == 0:
        raise ValueError("No data rows were found in the file (or sheet).")
    return df
